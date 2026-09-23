#!/usr/bin/env python3
"""Build the 3-slide multi-entity coding deck (editable PowerPoint) that mirrors public/coding-efficiency.html.

Reuses the shape/text helpers from the hub's demo-deck skill (read from ~/.claude/ottimate-hub-path.txt).
Usage: python3 deck/make_deck.py [out.pptx]
"""
import sys
from pathlib import Path

HUB = Path(Path.home(), ".claude/ottimate-hub-path.txt").read_text().strip()
sys.path.insert(0, f"{HUB}/skills/demo-deck/scripts")
from build_deck import rect, text  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

FONT, MONO = "Segoe UI", "Consolas"
hexc = lambda h: RGBColor.from_string(h)
SLATE9, SLATE8, SLATE6, SLATE5, SLATE4, SLATE2, SLATE1, SLATE0 = map(hexc, ["0F172A", "1E293B", "475569", "64748B", "94A3B8", "E2E8F0", "F1F5F9", "F8FAFC"])
BLUE6, BLUE4, BLUE1 = hexc("2563EB"), hexc("60A5FA"), hexc("EFF6FF")
EM6, EM5, EM3, EM1, EM8 = hexc("059669"), hexc("10B981"), hexc("6EE7B7"), hexc("ECFDF5"), hexc("065F46")
INDIGO, PURPLE, EMTOK, OTHER = hexc("A5B4FC"), hexc("C084FC"), hexc("34D399"), hexc("CBD5E1")
RED, RED1, AMBER, AMBER1 = hexc("B91C1C"), hexc("FEF2F2"), hexc("92400E"), hexc("FFFBEB")
WHITE = hexc("FFFFFF")
L, W = Inches(0.6), Inches(12.13)

# Worked numbers for slide 3: 4 fields per invoice, 15 seconds to map one field.
SECS, FIELDS = 15, 4
VOLUMES = [1000, 5000, 20000]


def t(s, x, y, w, h, txt, size, color, bold=False, align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP):
    text(s, x, y, w, h, [[(txt, size, color, bold, font)]], align=align, anchor=anchor, after=0)


def chip(s, x, y, w, h, label, fg, bg, line=None, size=12, bold=True):
    rect(s, x, y, w, h, fill=bg, line=line, rounded=True)
    t(s, x + Inches(.08), y, w - Inches(.16), h, label, size, fg, bold, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def slide(prs, eyebrow, title, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, prs.slide_width, prs.slide_height, fill=SLATE0)
    rect(s, 0, 0, prs.slide_width, Inches(1.2), fill=SLATE9)
    chip(s, L, Inches(.25), Inches(1.25), Inches(.32), "OTTIMATE", WHITE, BLUE6, size=11)
    t(s, L + Inches(1.4), Inches(.25), Inches(8), Inches(.32), eyebrow.upper(), 11, BLUE4, True, anchor=MSO_ANCHOR.MIDDLE)
    t(s, L, Inches(.62), W, Inches(.5), title, 26, WHITE, True)
    s.notes_slide.notes_text_frame.text = notes
    return s


def fields_slide(prs):
    s = slide(prs, "Multi-entity coding", "What your ERP needs on every invoice, and who fills it in",
              "Every ERP needs these fields before it will accept an invoice. Walk left to right.\n"
              "Typical AP tools can fill one or two: GL account from a vendor default, and sometimes entity if each "
              "entity gets its own inbox. Give them that credit, it makes the rest believable.\n"
              "Ottimate sets entity, location and department from where the invoice lands, because those values are "
              "fixed to the company, location and department in setup. Nobody maps them.\n"
              "Other fields (class, project, fund): fill by location or rule where the value is fixed. Confirm per "
              "customer before promising it.")
    label_w, gap = Inches(1.9), Inches(.12)
    col_w = int((W - label_w - 5 * gap) / 5)
    xs = [L + label_w + gap + i * (col_w + gap) for i in range(5)]
    toks = [("ENTITY", INDIGO), ("LOCATION", BLUE4), ("DEPARTMENT", PURPLE), ("GL ACCOUNT", EMTOK), ("OTHER FIELDS", OTHER)]
    y = Inches(1.6)
    for x, (tok, c) in zip(xs, toks):
        rect(s, x, y, col_w, Inches(.55), fill=SLATE8, rounded=True)
        t(s, x, y, col_w, Inches(.55), tok, 14, c, True, PP_ALIGN.CENTER, MONO, MSO_ANCHOR.MIDDLE)

    hand, dflt, auto, rule = (RED, RED1), (AMBER, AMBER1), (WHITE, EM5), (SLATE6, SLATE1)
    rows = [
        ("Typical AP tool", SLATE6, WHITE, SLATE2, [("By hand, or a separate inbox per entity", dflt), ("Mapped by hand", hand),
                                                     ("Mapped by hand", hand), ("Vendor default", dflt), ("By hand or rule", rule)]),
        ("Ottimate", EM6, EM1, EM3, [("Set by your structure", auto), ("Set by your structure", auto),
                                     ("Set by your structure", auto), ("Vendor default", dflt), ("By location or rule", rule)]),
    ]
    for r, (name, fg, bg, line, cells) in enumerate(rows):
        ry, rh = Inches(2.4 + r * 1.55), Inches(1.35)
        rect(s, L, ry, W, rh, fill=bg, line=line, rounded=True)
        t(s, L + Inches(.25), ry, label_w - Inches(.25), rh, name, 17, fg, True, anchor=MSO_ANCHOR.MIDDLE)
        for x, (label, (cfg, cbg)) in zip(xs, cells):
            chip(s, x + Inches(.1), ry + Inches(.3), col_w - Inches(.2), Inches(.75), label, cfg, cbg, size=13)

    rect(s, xs[0], Inches(5.45), xs[3] - xs[0] - gap, Inches(.08), fill=EM5)
    t(s, xs[0], Inches(5.6), xs[3] - xs[0] - gap, Inches(.4), "Ottimate handles these. Your team never maps them.", 15, EM6, True, PP_ALIGN.CENTER)
    rect(s, L, Inches(6.25), W, Inches(.8), fill=SLATE9, rounded=True)
    t(s, L + Inches(.3), Inches(6.25), W - Inches(.6), Inches(.8),
      "Typical tools fill one or two fields. Ottimate fills entity, location and department on every invoice.",
      17, WHITE, True, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def how_slide(prs):
    s = slide(prs, "How it works", "Set your structure once. Every invoice lands already coded.",
              "Left: the setup. Each level owns its part of the coding, so an invoice inherits all of it the moment "
              "it lands in a location.\n"
              "Right: one email address for everything. Ottimate places each invoice by ship-to address, customer "
              "number or ZIP code. Anything it cannot place goes to a To Be Sorted queue, so nothing is lost.\n"
              "Division of labor: managers see and approve only their own invoices; central AP sees every entity "
              "under one login.")
    colw = Inches(5.85)
    t(s, L, Inches(1.5), colw, Inches(.4), "Your coding lives on your structure", 18, SLATE8, True)
    levels = [("COMPANY", "Sets the entity", INDIGO), ("LOCATION", "Sets the location", BLUE4),
              ("DEPARTMENT", "Sets the department", PURPLE)]
    for i, (k, v, c) in enumerate(levels):
        y = Inches(2.05 + i * 1.0)
        rect(s, L + Inches(i * .35), y, colw - Inches(i * .35), Inches(.8), fill=WHITE, line=SLATE2, rounded=True)
        rect(s, L + Inches(i * .35), y, Inches(.12), Inches(.8), fill=c)
        t(s, L + Inches(i * .35 + .3), y + Inches(.1), Inches(3), Inches(.3), k, 11, SLATE5, True)
        t(s, L + Inches(i * .35 + .3), y + Inches(.38), Inches(4), Inches(.35), v, 16, SLATE8, True)
    y = Inches(5.1)
    rect(s, L, y, colw, Inches(1.25), fill=SLATE8, rounded=True)
    t(s, L + Inches(.3), y + Inches(.15), colw, Inches(.3), "INVOICE ARRIVES CODED", 11, SLATE4, True)
    text(s, L + Inches(.3), y + Inches(.5), colw - Inches(.6), Inches(.5),
         [[("ENTITY", 17, INDIGO, True, MONO), ("-", 17, SLATE5, True, MONO), ("LOCATION", 17, BLUE4, True, MONO),
           ("-", 17, SLATE5, True, MONO), ("DEPT", 17, PURPLE, True, MONO), ("-", 17, SLATE5, True, MONO),
           ("GL", 17, EMTOK, True, MONO)]], after=0)

    x = L + colw + Inches(.43)
    t(s, x, Inches(1.5), colw, Inches(.4), "One inbox, every team", 18, SLATE8, True)
    steps = [("One email address", "Vendors and staff send every invoice to the same place."),
             ("Placed automatically", "Ship-to address, customer number or ZIP code tell Ottimate where it belongs."),
             ("Arrives already coded", "Entity, location and department fill in the moment it lands."),
             ("Straight to the right manager", "Each manager approves only their own invoices.")]
    for i, (h, d) in enumerate(steps):
        y = Inches(2.05 + i * 1.1)
        rect(s, x, y, colw, Inches(.95), fill=EM1, line=EM3, rounded=True)
        chip(s, x + Inches(.2), y + Inches(.25), Inches(.45), Inches(.45), str(i + 1), WHITE, EM5, size=14)
        t(s, x + Inches(.85), y + Inches(.12), colw - Inches(1), Inches(.35), h, 16, SLATE8, True)
        t(s, x + Inches(.85), y + Inches(.5), colw - Inches(1), Inches(.4), d, 12, SLATE6)


def worth_slide(prs):
    s = slide(prs, "What it's worth", "Hours back every year, at your volume",
              f"The math: invoices x fields the typical tool leaves to your team x {SECS} seconds per field x 12 months.\n"
              f"Assumes {FIELDS} fields per invoice (entity, location, department, GL account). GL account is suggested "
              "by a vendor default in both tools, so it counts the same on both sides and is left out.\n"
              "Typical tool fills 1 field: your team maps 3 by hand there. Fills 2: your team maps 2.\n"
              "75% = 3 of 4 fields set by the customer's structure. 'Up to' because invoices split across locations, or "
              "needing an override, save less. Invoices with many separately coded lines save more.\n"
              "Live version with a volume slider: the coding-efficiency page in the config explainer.")
    rect(s, L, Inches(1.6), Inches(3.9), Inches(4.4), fill=WHITE, line=EM3, rounded=True)
    t(s, L, Inches(2.4), Inches(3.9), Inches(1.0), "Up to 75%", 48, EM6, True, PP_ALIGN.CENTER)
    t(s, L + Inches(.3), Inches(3.6), Inches(3.3), Inches(1.2), "of every invoice coded by your structure, not your team",
      16, EM8, True, PP_ALIGN.CENTER)

    x, colw = L + Inches(4.3), Inches(12.13 - 4.3)
    heads = ["Invoices per month", "Typical tool fills 1 field", "Typical tool fills 2 fields"]
    widths = [Inches(2.6), (colw - Inches(2.6)) / 2, (colw - Inches(2.6)) / 2]
    xs = [x, x + widths[0], x + widths[0] + widths[1]]
    for cx, cw, h in zip(xs, widths, heads):
        rect(s, cx, Inches(1.6), cw, Inches(.7), fill=SLATE8)
        t(s, cx, Inches(1.6), cw, Inches(.7), h, 14, WHITE, True, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for r, vol in enumerate(VOLUMES):
        y = Inches(2.3 + r * 1.05)
        vals = [f"{vol:,}"] + [f"{round(vol * hand * SECS / 3600 * 12):,} hrs / yr" for hand in (FIELDS - 1, FIELDS - 2)]
        for i, (cx, cw, v) in enumerate(zip(xs, widths, vals)):
            rect(s, cx, y, cw, Inches(1.05), fill=WHITE if r % 2 == 0 else SLATE1, line=SLATE2)
            t(s, cx, y, cw, Inches(1.05), v, 24 if i else 20, EM6 if i else SLATE8, True, PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)
    t(s, x, Inches(5.55), colw, Inches(.4), f"Assumes {FIELDS} fields per invoice and {SECS} seconds to map one field.",
      12, SLATE5, False, PP_ALIGN.CENTER)
    rect(s, L, Inches(6.25), W, Inches(.8), fill=SLATE9, rounded=True)
    t(s, L, Inches(6.25), W, Inches(.8), "Set it up once. Get the time back on every invoice, at every location.",
      17, WHITE, True, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def main(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    fields_slide(prs)
    how_slide(prs)
    worth_slide(prs)
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "multi-entity-coding.pptx")
